#!/usr/bin/env python3
"""Diff the member rows of two filled-in TBP Report of Eligibility workbooks.

Useful for reconciling two versions of the same report (e.g. one prepared by
this script and one prepared by hand). Reads the hidden ``TBPEligMembers``
sheet from each ``.xls`` file and writes an ``.xlsx`` workbook listing rows
that appear in only one of the two.

Usage:
    python compare_reports.py <a.xls> <b.xls> [-o discrepancies.xlsx]
"""

import argparse
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
import xlrd

SHEET = "TBPEligMembers"
MEMBER_COLS = list(range(11, 20))  # Fname(11) through Email(19)
HEADERS = [
    "Fname", "Mname", "Lname", "Class_Type", "ClassM", "ClassY",
    "Curriculum", "PresMem", "Email",
]

_SECTION_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_SECTION_FONT = Font(color="FFFFFF", bold=True, size=12)
_HEADER_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
_HEADER_FONT = Font(bold=True)


def read_entries(path: Path) -> list[tuple[str, ...]]:
    """Return the member rows of *path* as tuples of stripped strings."""
    wb = xlrd.open_workbook(str(path))
    ws = wb.sheet_by_name(SHEET)
    entries: list[tuple[str, ...]] = []
    for i in range(1, ws.nrows):
        vals = []
        for c in MEMBER_COLS:
            v = ws.cell_value(i, c)
            if isinstance(v, float) and v == int(v):
                v = int(v)
            vals.append(str(v).strip())
        if not vals[0]:  # skip rows with no Fname
            continue
        entries.append(tuple(vals))
    return entries


def only_in(a: Counter, b: Counter) -> list[tuple[str, ...]]:
    """Rows in multiset *a* that are not in *b* (duplicates respected)."""
    out: list[tuple[str, ...]] = []
    for entry, count in a.items():
        out.extend([entry] * max(count - b.get(entry, 0), 0))
    out.sort(key=lambda e: (e[2], e[0], e[1]))  # last, first, middle
    return out


def _styled_row(ws, values, fill, font) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.fill = fill
        cell.font = font


def _section(ws, title: str, rows: list[tuple[str, ...]]) -> None:
    _styled_row(ws, [title] + [""] * (len(HEADERS) - 1), _SECTION_FILL, _SECTION_FONT)
    _styled_row(ws, HEADERS, _HEADER_FILL, _HEADER_FONT)
    for entry in rows:
        ws.append(list(entry))


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Diff the member rows of two TBP Report of Eligibility .xls files."
    )
    parser.add_argument("a", type=Path, help="First report (.xls).")
    parser.add_argument("b", type=Path, help="Second report (.xls).")
    parser.add_argument(
        "-o", "--output", type=Path, default=Path("discrepancies.xlsx"),
        help="Output workbook path (default: discrepancies.xlsx).",
    )
    args = parser.parse_args()
    for p in (args.a, args.b):
        if not p.exists():
            print(f"Error: {p} not found.", file=sys.stderr)
            sys.exit(1)

    a_entries = read_entries(args.a)
    b_entries = read_entries(args.b)
    a_counts, b_counts = Counter(a_entries), Counter(b_entries)
    only_a = only_in(a_counts, b_counts)
    only_b = only_in(b_counts, a_counts)

    print(f"{args.a.name}: {len(a_entries)} entries")
    print(f"{args.b.name}: {len(b_entries)} entries")
    print(f"Only in {args.a.name}: {len(only_a)}")
    print(f"Only in {args.b.name}: {len(only_b)}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Discrepancies"
    _section(ws, f"Only in {args.a.name} ({len(only_a)} entries)", only_a)
    ws.append([])
    ws.append([])
    _section(ws, f"Only in {args.b.name} ({len(only_b)} entries)", only_b)

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    wb.save(args.output)
    print(f"\nWrote {args.output}")


if __name__ == "__main__":
    main()
