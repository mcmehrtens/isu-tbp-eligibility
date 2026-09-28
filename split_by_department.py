#!/usr/bin/env python3
"""Split eligible non-members into per-department lists.

Reads the ISU eligibility workbook (optionally password-protected), keeps the
eligible students who are not already TBP members, and assigns each to an
engineering department by major. Produces:

1. An Excel workbook with one tab per department (name, Net-ID, major),
   encrypted with the input password.
2. One CSV of Net-IDs per department.
3. Updates ``Major-Department Reference.xlsx`` (tracked in this repo), which
   lists every major seen so far and the department it belongs to. Majors
   from the current (and any ``--prior``) workbook are merged into the
   existing file, so its history accumulates across semesters.

Usage:
    python split_by_department.py <semester> <input.xlsx> \
        [--prior SEMESTER=PATH ...] [-o OUTDIR] [--reference PATH] [--no-encrypt]

The password is read from the ``TBP_PASSWORD`` environment variable if set;
otherwise you are prompted for it when an input is encrypted.
"""

import argparse
import csv
import getpass
import io
import os
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from msoffcrypto.format.ooxml import OOXMLFile
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from prep_eligibility_data import (
    _HEADER_FILL,
    _HEADER_FONT,
    decrypt_workbook,
    parse_current_semester,
    read_sheet,
)

# =============================================================================
# CONFIGURATION — Edit these values as needed for future semesters.
# =============================================================================

# Sheets in the input workbook that hold eligible students.
ELIGIBLE_SHEETS = ["Juniors", "Seniors", "Masters", "PhD"]

# Sheets that list current members (name differs between semesters). Used as
# a safety net on top of the TBP columns: anyone listed here is excluded.
MEMBER_SHEETS = ["Current TBP Members", "Members"]

# Column headers in the input spreadsheet.
COL_FIRST = "First Name"
COL_LAST = "Last Name"
COL_PROGRAM = "Program of Study"
COL_EMAIL = "PRIMARY_INSTITUTIONAL_EMAIL_ADDRESS"
COL_TBP_GRAD = "Tau_Beta_Pi_Grad"
COL_TBP_UNDERGRAD = "Tau_Beta_Pi_Undergrad"

# Departments, in the order their tabs and CSVs are written.
DEPARTMENTS: dict[str, str] = {
    "AERE": "Aerospace Engineering",
    "ABE": "Agricultural and Biosystems Engineering",
    "CBE": "Chemical and Biological Engineering",
    "CCEE": "Civil, Construction and Environmental Engineering",
    "ECpE": "Electrical and Computer Engineering",
    "IMSE": "Industrial and Manufacturing Systems Engineering",
    "MSE": "Materials Science and Engineering",
    "ME": "Mechanical Engineering",
}

# Program name (the "Program of Study" text before the first comma, i.e.
# without the degree) -> department. If ISU adds or renames a program, the
# script stops and lists it; add an entry here.
PROGRAM_DEPARTMENTS: dict[str, str] = {
    "Aerospace Engineering": "AERE",
    "Engineering Mechanics": "AERE",
    "Agricultural Engineering": "ABE",
    "Agricultural and Biosystems Engineering": "ABE",
    "Biological Systems Engineering": "ABE",
    "Industrial and Agricultural Technology": "ABE",
    "Biomedical Engineering": "CBE",
    "Chemical Engineering": "CBE",
    "Civil Engineering": "CCEE",
    "Construction Engineering": "CCEE",
    "Environmental Engineering": "CCEE",
    "Computer Engineering": "ECpE",
    "Cyber Security Engineering": "ECpE",
    "Electrical Engineering": "ECpE",
    "Software Engineering": "ECpE",
    "Engineering Management": "IMSE",
    "Industrial Engineering": "IMSE",
    "Systems Engineering": "IMSE",
    "Materials Engineering": "MSE",
    "Materials Science and Engineering": "MSE",
    "Energy Systems Engineering": "ME",
    "Mechanical Engineering": "ME",
}

# Why non-obvious programs map where they do (shown in the reference workbook).
DEPARTMENT_NOTES: dict[str, str] = {
    "Engineering Mechanics": (
        "Graduate program administered by Aerospace Engineering. "
        "https://catalog.iastate.edu/collegeofengineering/engineeringmechanics/"
    ),
    "Industrial and Agricultural Technology": (
        "Graduate program housed in ABE. "
        "https://www.abe.iastate.edu/graduate-students/industrial-and-agricultural-technology/"
    ),
    "Agricultural and Biosystems Engineering": (
        "ABE is shared by Engineering and CALS; the (ABE E)/(ABE A) suffix is "
        "the college, not a different department."
    ),
    "Biomedical Engineering": (
        "Interdepartmental B.S.; the professor in charge (Ian Schneider) and "
        "BME advising are in CBE (Sweeney Hall). Placed in CBE, as in Spring 2026. "
        "https://catalog.iastate.edu/collegeofengineering/biomedicalengineering/"
    ),
    "Environmental Engineering": (
        "Administered by CCEE. "
        "https://catalog.iastate.edu/collegeofengineering/environmentalengineering/"
    ),
    "Cyber Security Engineering": (
        "Administered by ECpE. "
        "https://catalog.iastate.edu/collegeofengineering/cybersecurityengineering/"
    ),
    "Software Engineering": (
        "Joint between Engineering (ECpE) and LAS (Computer Science); ECpE is "
        "the engineering home. "
        "https://catalog.iastate.edu/collegeofengineering/softwareengineering/"
    ),
    "Engineering Management": (
        "Administered by IMSE. "
        "https://catalog.iastate.edu/interdisciplinaryprograms/engineeringmanagement/"
    ),
    "Systems Engineering": (
        "M.Engr. offered through IMSE. "
        "https://www.imse.iastate.edu/graduate-program/adminfo/admsyse/"
    ),
    "Energy Systems Engineering": (
        "Interdisciplinary M.Engr. offered by ME. "
        "https://www.me.iastate.edu/graduate-programs/meng-degree-in-energy-systems-engineering/"
    ),
}

# Output file names; {semester} becomes e.g. "Fall 2026".
DEPT_WORKBOOK_NAME = "{semester} Non-Member Lists - Separated by Department.xlsx"
CSV_DIR_NAME = "Department Net-IDs"
CSV_NAME = "{semester} Net-IDs - {dept}.csv"

# Major -> department reference. Contains no student data, so it is tracked
# in the repo (see the exception in .gitignore) rather than written to OUTDIR.
REFERENCE_WORKBOOK = Path(__file__).resolve().parent / "Major-Department Reference.xlsx"

# =============================================================================
# END CONFIGURATION
# =============================================================================


# -- Helpers -------------------------------------------------------------------


def semester_label(value: str) -> str:
    """Turn a CLI semester like ``F2026`` into ``"Fall 2026"``."""
    year, sem = parse_current_semester(value)
    return f"{sem} {year}"


def _label_key(label: str) -> tuple[int, int]:
    """Sort key for a label like ``"Fall 2026"`` (chronological)."""
    sem, year = label.split()
    return int(year), 0 if sem == "Spring" else 1


def split_program(program: str) -> tuple[str, str]:
    """Split a "Program of Study" value into ``(name, degree)``.

    ``"Civil Engineering, M.Engr (Distance)"`` -> ``("Civil Engineering",
    "M.Engr (Distance)")``.
    """
    name, _, degree = program.partition(",")
    return name.strip(), degree.strip()


def net_id(email: str) -> str:
    """Return the Net-ID (the part of the email before ``@``), lowercased."""
    return email.split("@", 1)[0].strip().lower()


class WorkbookOpener:
    """Open workbooks, decrypting with a password obtained at most once."""

    def __init__(self) -> None:
        self.password: str | None = os.environ.get("TBP_PASSWORD")

    def get_password(self) -> str:
        """Return the password, prompting for it if not already known."""
        if self.password is None:
            self.password = getpass.getpass("Enter spreadsheet password: ")
        return self.password

    def open(self, path: Path) -> openpyxl.Workbook:
        """Open *path* read-only, decrypting it if necessary.

        Exits with an error message if the file cannot be opened.
        """
        if not path.exists():
            print(f"Error: {path} not found.", file=sys.stderr)
            sys.exit(1)
        try:
            return openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception:
            pass
        try:
            return decrypt_workbook(path, self.get_password())
        except Exception as exc:
            print(f"Error: Could not decrypt {path} — {exc}", file=sys.stderr)
            sys.exit(1)


def read_programs(wb: openpyxl.Workbook) -> set[str]:
    """Return every non-empty "Program of Study" value on any sheet of *wb*."""
    programs: set[str] = set()
    for name in wb.sheetnames:
        for row in read_sheet(wb, name):
            if program := row.get(COL_PROGRAM, ""):
                programs.add(program)
    return programs


def check_programs(programs: set[str]) -> None:
    """Exit with an error if any program has no department mapping."""
    unknown = sorted(p for p in programs if split_program(p)[0] not in PROGRAM_DEPARTMENTS)
    if unknown:
        print(
            "Error: these programs are not in PROGRAM_DEPARTMENTS. Add them to the\n"
            "CONFIGURATION block and re-run:",
            file=sys.stderr,
        )
        for p in unknown:
            print(f"    - {p!r}", file=sys.stderr)
        sys.exit(1)


# -- Core logic ----------------------------------------------------------------


def collect_non_members(wb: openpyxl.Workbook) -> list[dict[str, str]]:
    """Return one record per eligible non-member, merged by Net-ID.

    A student is a non-member when both TBP columns are blank and their email
    is not on any member sheet. Students listed on more than one eligible
    sheet (e.g. both Masters and PhD) are merged, with their majors joined.

    Parameters
    ----------
    wb : openpyxl.Workbook
        The current semester's eligibility workbook.

    Returns
    -------
    list[dict[str, str]]
        Records with keys ``last``, ``first``, ``netid``, ``majors`` (list),
        and ``dept``.
    """
    member_emails: set[str] = set()
    for name in MEMBER_SHEETS:
        if name in wb.sheetnames:
            member_emails |= {
                r[COL_EMAIL].lower() for r in read_sheet(wb, name) if r.get(COL_EMAIL)
            }

    people: dict[str, dict] = {}
    excluded_by_member_sheet = 0
    for sheet in ELIGIBLE_SHEETS:
        for row in read_sheet(wb, sheet):
            email = row.get(COL_EMAIL, "")
            if not email:
                continue  # blank row
            if row.get(COL_TBP_GRAD) or row.get(COL_TBP_UNDERGRAD):
                continue
            if email.lower() in member_emails:
                excluded_by_member_sheet += 1
                continue
            program = row[COL_PROGRAM]
            dept = PROGRAM_DEPARTMENTS[split_program(program)[0]]
            nid = net_id(email)
            if nid in people:
                person = people[nid]
                if program not in person["majors"]:
                    person["majors"].append(program)
                if person["dept"] != dept:
                    print(
                        f"Warning: {nid} has majors in {person['dept']} and {dept}; "
                        f"listed under {person['dept']} only.",
                        file=sys.stderr,
                    )
                continue
            people[nid] = {
                "last": row.get(COL_LAST, ""),
                "first": row.get(COL_FIRST, ""),
                "netid": nid,
                "majors": [program],
                "dept": dept,
            }

    if excluded_by_member_sheet:
        print(
            f"Excluded {excluded_by_member_sheet} row(s) with blank TBP columns "
            "that appear on a member sheet."
        )
    return sorted(
        people.values(),
        key=lambda p: (p["last"].casefold(), p["first"].casefold(), p["netid"]),
    )


# -- Output --------------------------------------------------------------------


def _write_table(ws, headers: list[str], rows: list[list[str]]) -> None:
    """Write a header row and data rows with the repo's header styling."""
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, header in enumerate(headers, start=1):
        width = max([len(header)] + [len(str(r[col_idx - 1])) for r in rows])
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = min(width + 2, 80)


def write_department_workbook(
    path: Path, by_dept: dict[str, list[dict]], password: str | None
) -> None:
    """Write one tab per department, encrypting the file if *password* is set."""
    wb = openpyxl.Workbook()
    wb.remove(wb.worksheets[0])
    for dept, people in by_dept.items():
        ws = wb.create_sheet(dept)
        _write_table(
            ws,
            ["Last Name", "First Name", "Net-ID", "Major"],
            [[p["last"], p["first"], p["netid"], "; ".join(p["majors"])] for p in people],
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    with open(path, "wb") as f:
        if password:
            OOXMLFile(buf).encrypt(password, f)
        else:
            f.write(buf.getvalue())


def write_csvs(directory: Path, semester: str, by_dept: dict[str, list[dict]]) -> None:
    """Write one header-less CSV of Net-IDs per department."""
    directory.mkdir(parents=True, exist_ok=True)
    for dept, people in by_dept.items():
        path = directory / CSV_NAME.format(semester=semester, dept=dept)
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f, lineterminator="\n")
            writer.writerows([p["netid"]] for p in people)


def read_reference_workbook(path: Path) -> dict[str, list[str]]:
    """Read Program of Study -> "Seen In" semesters from an existing reference.

    Returns an empty dict if *path* does not exist.
    """
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, read_only=True)
    seen: dict[str, list[str]] = {}
    for row in read_sheet(wb, wb.sheetnames[0]):
        if program := row.get("Program of Study", ""):
            seen[program] = [s.strip() for s in row.get("Seen In", "").split(",") if s.strip()]
    wb.close()
    return seen


def write_reference_workbook(path: Path, seen: dict[str, list[str]]) -> None:
    """Write the major -> department reference sheet.

    Parameters
    ----------
    path : Path
        Output ``.xlsx`` path.
    seen : dict[str, list[str]]
        Program of Study -> semester labels it appeared in (chronological).
    """
    dept_order = list(DEPARTMENTS)
    rows = []
    for program, semesters in seen.items():
        name, degree = split_program(program)
        dept = PROGRAM_DEPARTMENTS[name]
        rows.append([
            dept, DEPARTMENTS[dept], program, degree,
            ", ".join(semesters), DEPARTMENT_NOTES.get(name, ""),
        ])
    rows.sort(key=lambda r: (dept_order.index(r[0]), r[2].casefold()))

    wb = openpyxl.Workbook()
    wb.remove(wb.worksheets[0])
    ws = wb.create_sheet("Majors")
    _write_table(
        ws,
        ["Dept", "Department Name", "Program of Study", "Degree", "Seen In", "Notes"],
        rows,
    )
    ws.column_dimensions["F"].width = 70
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


# -- Main ----------------------------------------------------------------------


def _prior_arg(value: str) -> tuple[str, Path]:
    """Parse ``--prior S2026=path/to/lists.xlsx``."""
    sem, sep, path = value.partition("=")
    if not sep:
        raise argparse.ArgumentTypeError(f"Expected SEMESTER=PATH, got {value!r}.")
    return semester_label(sem), Path(path)


def main() -> None:
    """Parse arguments, build the department lists, and write all outputs."""
    parser = argparse.ArgumentParser(
        description="Split eligible non-members into per-department lists."
    )
    parser.add_argument(
        "semester", type=semester_label, help="Current semester (e.g. F2026)."
    )
    parser.add_argument("input", type=Path, help="Current ISU eligibility .xlsx.")
    parser.add_argument(
        "--prior",
        type=_prior_arg,
        action="append",
        default=[],
        metavar="SEMESTER=PATH",
        help="An earlier semester's workbook, used only for the majors reference "
        "(repeatable), e.g. S2026=data/Spring2026TBPLists.xlsx.",
    )
    parser.add_argument(
        "-o", "--outdir", type=Path, default=Path("data"),
        help="Output directory (default: data/).",
    )
    parser.add_argument(
        "--reference", type=Path, default=REFERENCE_WORKBOOK,
        help="Major -> department reference workbook to update "
        "(default: the one in this repo).",
    )
    parser.add_argument(
        "--no-encrypt", action="store_true",
        help="Do not password-protect the department workbook.",
    )
    args = parser.parse_args()
    semester: str = args.semester

    opener = WorkbookOpener()

    # -- Programs seen, for validation and the reference sheet -----------------
    # Start from the existing reference so earlier semesters are kept.
    seen = read_reference_workbook(args.reference)
    for label, path in [*args.prior, (semester, args.input)]:
        wb = opener.open(path)
        for program in read_programs(wb):
            labels = seen.setdefault(program, [])
            if label not in labels:
                labels.append(label)
        wb.close()
    for labels in seen.values():
        labels.sort(key=_label_key)
    check_programs(set(seen))

    # -- Non-members by department ---------------------------------------------
    wb = opener.open(args.input)
    people = collect_non_members(wb)
    wb.close()

    by_dept: dict[str, list[dict]] = {d: [] for d in DEPARTMENTS}
    for p in people:
        by_dept[p["dept"]].append(p)

    # -- Write outputs ---------------------------------------------------------
    args.outdir.mkdir(parents=True, exist_ok=True)
    dept_path = args.outdir / DEPT_WORKBOOK_NAME.format(semester=semester)
    password = None if args.no_encrypt else opener.password
    write_department_workbook(dept_path, by_dept, password)
    write_csvs(args.outdir / CSV_DIR_NAME, semester, by_dept)
    write_reference_workbook(args.reference, seen)

    # -- Summary ---------------------------------------------------------------
    print(f"\n{len(people)} non-members across {len(DEPARTMENTS)} departments:")
    for dept, dept_people in by_dept.items():
        print(f"  {dept:<5} {len(dept_people):>4}")
        majors = Counter(m for p in dept_people for m in p["majors"])
        for major, count in sorted(majors.items()):
            print(f"          {count:>4}  {major}")
    merged = [p for p in people if len(p["majors"]) > 1]
    for p in merged:
        print(f"Merged duplicate: {p['first']} {p['last']} ({p['netid']}): "
              f"{'; '.join(p['majors'])}")
    print(f"\nWrote {dept_path}" + (" (encrypted)" if password else ""))
    print(f"Wrote {args.outdir / CSV_DIR_NAME}/ ({len(DEPARTMENTS)} CSVs)")
    print(f"Updated {args.reference} ({len(seen)} programs)")


if __name__ == "__main__":
    main()
